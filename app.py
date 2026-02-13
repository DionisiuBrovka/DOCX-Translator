"""
DOCX Translator — Перевод Word-документов с сохранением форматирования.
Использует Groq API для быстрого перевода через LLM.
"""

import os
import json
import time
import uuid
import logging
from pathlib import Path
from copy import deepcopy
from typing import Optional

from flask import Flask, render_template, request, send_file, jsonify, session
from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from groq import Groq

# ---------------------------------------------------------------------------
# Config
# ---------------------------------------------------------------------------

app = Flask(__name__)
app.secret_key = os.environ.get("SECRET_KEY", "docx-translator-secret-key-change-me")

UPLOAD_FOLDER = Path("uploads")
OUTPUT_FOLDER = Path("outputs")
UPLOAD_FOLDER.mkdir(exist_ok=True)
OUTPUT_FOLDER.mkdir(exist_ok=True)

GROQ_API_KEY = os.environ.get("GROQ_API_KEY", "")
GROQ_MODEL = os.environ.get("GROQ_MODEL", "llama-3.3-70b-versatile")

logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)

# ---------------------------------------------------------------------------
# Translation Engine
# ---------------------------------------------------------------------------

class TranslationEngine:
    """Переводчик текста через Groq API с сохранением структуры ранов."""

    def __init__(self, api_key: str, model: str = GROQ_MODEL,
                 source_lang: str = "Russian", target_lang: str = "English"):
        self.client = Groq(api_key=api_key)
        self.model = model
        self.source_lang = source_lang
        self.target_lang = target_lang
        self.request_count = 0
        self.total_tokens = 0

    def translate_runs(self, runs_text: list[str]) -> list[str]:
        """
        Переводит список текстовых фрагментов (ранов), сохраняя их количество
        и соответствие. Это ключ к сохранению inline-форматирования.

        Вход:  ["Привет, ", "мир!"]
        Выход: ["Hello, ", "world!"]
        """
        # Пропускаем пустые параграфы
        combined = "".join(runs_text).strip()
        if not combined:
            return runs_text

        # Если один ран — простой перевод
        if len(runs_text) == 1:
            translated = self._translate_simple(runs_text[0])
            return [translated]

        # Несколько ранов — структурный перевод
        return self._translate_structured(runs_text)

    def _translate_simple(self, text: str) -> str:
        """Простой перевод одного фрагмента текста."""
        if not text.strip():
            return text

        # Сохраняем ведущие/завершающие пробелы
        leading = text[:len(text) - len(text.lstrip())]
        trailing = text[len(text.rstrip()):]

        response = self.client.chat.completions.create(
            model=self.model,
            messages=[
                {
                    "role": "system",
                    "content": (
                        f"You are a professional translator from {self.source_lang} "
                        f"to {self.target_lang}. Translate the user's text accurately "
                        f"and naturally. Return ONLY the translated text, nothing else. "
                        f"Do not add explanations, notes, or quotes around the text. "
                        f"Preserve any leading/trailing whitespace exactly."
                    )
                },
                {"role": "user", "content": text.strip()}
            ],
            temperature=0.3,
            max_tokens=4096,
        )
        self.request_count += 1
        self.total_tokens += response.usage.total_tokens if response.usage else 0

        result = response.choices[0].message.content.strip()
        return leading + result + trailing

    def _translate_structured(self, runs_text: list[str]) -> list[str]:
        """
        Структурный перевод: сохраняем количество и порядок фрагментов.
        Отправляем как JSON-массив, получаем JSON-массив обратно.
        """
        # Фильтруем: запоминаем какие раны пустые
        indexed = [(i, t) for i, t in enumerate(runs_text)]
        non_empty = [(i, t) for i, t in indexed if t.strip()]

        if not non_empty:
            return runs_text

        fragments = [t for _, t in non_empty]
        fragments_json = json.dumps(fragments, ensure_ascii=False)

        response = self.client.chat.completions.create(
            model=self.model,
            messages=[
                {
                    "role": "system",
                    "content": (
                        f"You are a professional translator from {self.source_lang} "
                        f"to {self.target_lang}.\n\n"
                        f"TASK: Translate a JSON array of text fragments. These fragments "
                        f"form a single paragraph when concatenated. Each fragment has "
                        f"its own formatting (bold, italic, etc.) that must be preserved "
                        f"through maintaining the array structure.\n\n"
                        f"RULES:\n"
                        f"1. Return a valid JSON array with EXACTLY {len(fragments)} elements\n"
                        f"2. Each element corresponds to the same-index input element\n"
                        f"3. Preserve leading/trailing spaces in each element\n"
                        f"4. The concatenation of all elements must read as natural "
                        f"{self.target_lang}\n"
                        f"5. Return ONLY the JSON array, no explanation or markdown\n"
                        f"6. If a fragment is just whitespace or punctuation, keep it as-is"
                    )
                },
                {"role": "user", "content": fragments_json}
            ],
            temperature=0.2,
            max_tokens=4096,
        )
        self.request_count += 1
        self.total_tokens += response.usage.total_tokens if response.usage else 0

        raw = response.choices[0].message.content.strip()

        # Парсим JSON-ответ
        try:
            # Убираем возможные markdown-обёртки
            if raw.startswith("```"):
                raw = raw.split("\n", 1)[1] if "\n" in raw else raw[3:]
                if raw.endswith("```"):
                    raw = raw[:-3]
                raw = raw.strip()

            translated = json.loads(raw)

            if isinstance(translated, list) and len(translated) == len(fragments):
                # Собираем результат, подставляя переводы на места непустых ранов
                result = list(runs_text)  # копия
                for (orig_idx, _), trans in zip(non_empty, translated):
                    result[orig_idx] = str(trans)
                return result
        except (json.JSONDecodeError, ValueError) as e:
            logger.warning(f"Failed to parse structured response, falling back: {e}")

        # Fallback: переводим параграф целиком, кладём в первый ран
        full_text = "".join(runs_text)
        translated_full = self._translate_simple(full_text)
        result = [""] * len(runs_text)
        result[0] = translated_full
        return result


# ---------------------------------------------------------------------------
# Document Processor
# ---------------------------------------------------------------------------

class DocxTranslator:
    """Обходит .docx документ и переводит весь текст, сохраняя форматирование."""

    def __init__(self, engine: TranslationEngine):
        self.engine = engine
        self.stats = {
            "paragraphs": 0,
            "table_cells": 0,
            "headers_footers": 0,
            "skipped_empty": 0,
        }

    def translate_document(self, input_path: str, output_path: str,
                           progress_callback=None) -> dict:
        """Основной метод: читает, переводит, сохраняет."""
        doc = Document(input_path)

        # Считаем общее количество элементов для прогресса
        total = self._count_elements(doc)
        current = 0

        def update_progress(increment=1):
            nonlocal current
            current += increment
            if progress_callback and total > 0:
                progress_callback(min(current / total * 100, 100))

        # 1. Параграфы тела документа
        logger.info("Translating body paragraphs...")
        for para in doc.paragraphs:
            self._translate_paragraph(para)
            self.stats["paragraphs"] += 1
            update_progress()

        # 2. Таблицы
        logger.info("Translating tables...")
        for table in doc.tables:
            self._translate_table(table, update_progress)

        # 3. Хедеры и футеры
        logger.info("Translating headers and footers...")
        for section in doc.sections:
            for header in [section.header, section.first_page_header,
                           section.even_page_header]:
                if header and header.is_linked_to_previous is False:
                    for para in header.paragraphs:
                        self._translate_paragraph(para)
                        self.stats["headers_footers"] += 1
                        update_progress()
                    for table in header.tables:
                        self._translate_table(table, update_progress)

            for footer in [section.footer, section.first_page_footer,
                           section.even_page_footer]:
                if footer and footer.is_linked_to_previous is False:
                    for para in footer.paragraphs:
                        self._translate_paragraph(para)
                        self.stats["headers_footers"] += 1
                        update_progress()
                    for table in footer.tables:
                        self._translate_table(table, update_progress)

        # Сохраняем (картинки и прочие OLE-объекты сохраняются автоматически)
        doc.save(output_path)
        logger.info(f"Translation complete: {output_path}")

        return {
            **self.stats,
            "api_requests": self.engine.request_count,
            "total_tokens": self.engine.total_tokens,
        }

    def _translate_paragraph(self, para):
        """Переводит параграф, работая на уровне ранов."""
        runs = para.runs
        if not runs:
            self.stats["skipped_empty"] += 1
            return

        runs_text = [run.text for run in runs]
        combined = "".join(runs_text).strip()

        if not combined:
            self.stats["skipped_empty"] += 1
            return

        # Пропускаем, если текст не содержит кириллицу (вероятно, уже на англ.)
        if not any('\u0400' <= c <= '\u04FF' for c in combined):
            return

        translated = self.engine.translate_runs(runs_text)

        for run, new_text in zip(runs, translated):
            run.text = new_text

    def _translate_table(self, table, update_progress):
        """Переводит все ячейки таблицы."""
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    self._translate_paragraph(para)
                    self.stats["table_cells"] += 1
                    update_progress()
                # Вложенные таблицы
                for nested_table in cell.tables:
                    self._translate_table(nested_table, update_progress)

    def _count_elements(self, doc) -> int:
        """Считает общее количество переводимых элементов."""
        count = len(doc.paragraphs)
        for table in doc.tables:
            count += self._count_table_elements(table)
        for section in doc.sections:
            for hf in [section.header, section.footer,
                       section.first_page_header, section.first_page_footer,
                       section.even_page_header, section.even_page_footer]:
                if hf:
                    count += len(hf.paragraphs)
                    for table in hf.tables:
                        count += self._count_table_elements(table)
        return count

    def _count_table_elements(self, table) -> int:
        count = 0
        for row in table.rows:
            for cell in row.cells:
                count += len(cell.paragraphs)
                for nested in cell.tables:
                    count += self._count_table_elements(nested)
        return count


# ---------------------------------------------------------------------------
# Flask Routes
# ---------------------------------------------------------------------------

# Храним прогресс в памяти (для продакшена использовать Redis/БД)
translation_progress: dict[str, dict] = {}


@app.route("/")
def index():
    return render_template("index.html", groq_model=GROQ_MODEL)


@app.route("/translate", methods=["POST"])
def translate():
    """Загрузка и перевод документа."""
    # Проверяем API-ключ
    api_key = request.form.get("api_key", "").strip() or GROQ_API_KEY
    if not api_key:
        return jsonify({"error": "Groq API key is required"}), 400

    # Проверяем файл
    file = request.files.get("file")
    if not file or not file.filename:
        return jsonify({"error": "No file uploaded"}), 400

    if not file.filename.lower().endswith(".docx"):
        return jsonify({"error": "Only .docx files are supported"}), 400

    # Параметры
    model = request.form.get("model", GROQ_MODEL).strip()
    source_lang = request.form.get("source_lang", "Russian").strip()
    target_lang = request.form.get("target_lang", "English").strip()

    # Сохраняем загруженный файл
    task_id = str(uuid.uuid4())[:8]
    original_name = Path(file.filename).stem
    input_path = UPLOAD_FOLDER / f"{task_id}_{file.filename}"
    output_name = f"{original_name}_translated.docx"
    output_path = OUTPUT_FOLDER / f"{task_id}_{output_name}"

    file.save(str(input_path))

    # Инициализируем прогресс
    translation_progress[task_id] = {
        "progress": 0,
        "status": "processing",
        "output_name": output_name,
        "output_path": str(output_path),
    }

    # Переводим (синхронно; для продакшена — Celery/RQ)
    try:
        engine = TranslationEngine(
            api_key=api_key,
            model=model,
            source_lang=source_lang,
            target_lang=target_lang,
        )
        translator = DocxTranslator(engine)

        def on_progress(pct):
            translation_progress[task_id]["progress"] = round(pct, 1)

        stats = translator.translate_document(
            str(input_path), str(output_path), progress_callback=on_progress
        )

        translation_progress[task_id].update({
            "status": "done",
            "progress": 100,
            "stats": stats,
        })

        return jsonify({
            "task_id": task_id,
            "status": "done",
            "stats": stats,
            "download_url": f"/download/{task_id}",
        })

    except Exception as e:
        logger.exception("Translation failed")
        translation_progress[task_id].update({
            "status": "error",
            "error": str(e),
        })
        return jsonify({"error": str(e)}), 500

    finally:
        # Удаляем загруженный файл
        try:
            input_path.unlink()
        except OSError:
            pass


@app.route("/progress/<task_id>")
def progress(task_id):
    """Проверка прогресса перевода."""
    info = translation_progress.get(task_id)
    if not info:
        return jsonify({"error": "Task not found"}), 404
    return jsonify(info)


@app.route("/download/<task_id>")
def download(task_id):
    """Скачивание переведённого документа."""
    info = translation_progress.get(task_id)
    if not info or info["status"] != "done":
        return jsonify({"error": "File not ready"}), 404

    output_path = info["output_path"]
    if not Path(output_path).exists():
        return jsonify({"error": "File not found"}), 404

    return send_file(
        output_path,
        as_attachment=True,
        download_name=info["output_name"],
        mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )


# ---------------------------------------------------------------------------
# Run
# ---------------------------------------------------------------------------

if __name__ == "__main__":
    port = int(os.environ.get("PORT", 5000))
    debug = os.environ.get("FLASK_DEBUG", "0") == "1"
    app.run(host="0.0.0.0", port=port, debug=debug)
